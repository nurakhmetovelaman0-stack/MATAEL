
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[2]
REQUIREMENTS = ROOT / "docs" / "MATAEL_REQUIREMENTS.md"
OUTPUT = ROOT / "MATAEL_Исходный_код_и_требования.docx"

SOURCE_FILES = [
    ROOT / "app/src/main/java/com/matael/app/MainActivity.java",
    ROOT / "app/src/main/java/com/matael/app/MataelNote.java",
    ROOT / "app/src/main/java/com/matael/app/MataelSettings.java.",
]


def main():
    if not REQUIREMENTS.exists():
        raise FileNotFoundError(
            f"Не найден файл требований: {REQUIREMENTS}"
        )

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "MATAEL — исходный код и требования"
    )
    run.bold = True
    run.font.size = Pt(18)

    document.add_paragraph(
        "Документ сформирован из исходников GitHub."
    )

    document.add_heading(
        "Требования и план разработки", level=1
    )

    for line in REQUIREMENTS.read_text(
        encoding="utf-8"
    ).splitlines():
        text = line.strip()

        if not text:
            continue

        if text.startswith("# "):
            continue
        elif text.startswith("## "):
            document.add_heading(text[3:], level=2)
        elif text.startswith("- "):
            document.add_paragraph(
                text[2:], style="List Bullet"
            )
        else:
            document.add_paragraph(text)

    for source in SOURCE_FILES:
        if not source.exists():
            document.add_paragraph(
                f"Исходный файл не найден: {source.relative_to(ROOT)}"
            )
            continue

        document.add_page_break()
        document.add_heading(
            f"Исходный код: {source.relative_to(ROOT)}",
            level=1,
        )

        for line in source.read_text(
            encoding="utf-8"
        ).splitlines():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0

            code = paragraph.add_run(line if line else " ")
            code.font.name = "Consolas"
            code.font.size = Pt(7)

    document.save(OUTPUT)
    print(f"Word-документ создан: {OUTPUT}")


if __name__ == "__main__":
    main()
